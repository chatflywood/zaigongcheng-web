# -*- coding: utf-8 -*-
"""
analysis 四类预警导出路由集成测试。

覆盖：
  - 表头含「施工单位」
  - item 自带 constructionUnit 时直接写入
  - 旧预警 JSON 无该字段时，从 raw_data 按工程编码回查
  - Content-Disposition 文件名含导出日期后缀
"""
import io
import json
import re
import sys
from datetime import datetime
from pathlib import Path
from urllib.parse import unquote

import httpx
import openpyxl
import pytest
from fastapi import FastAPI
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

sys.path.insert(0, str(Path(__file__).parent.parent))

from models import Base, ZaigongRecord  # noqa: E402
import routers.analysis as analysis_mod  # noqa: E402


@pytest.fixture
def test_db():
    engine = create_engine(
        "sqlite:///:memory:",
        echo=False,
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    Session = sessionmaker(bind=engine, expire_on_commit=False)
    orig_get_db = analysis_mod.get_db
    analysis_mod.get_db = lambda: Session()
    try:
        yield Session
    finally:
        analysis_mod.get_db = orig_get_db
        Base.metadata.drop_all(engine)
        engine.dispose()


def build_client():
    app = FastAPI()
    app.include_router(analysis_mod.router, prefix="/api/zaigong")
    transport = httpx.ASGITransport(app=app)
    return httpx.AsyncClient(transport=transport, base_url="http://testserver")


def _seed_record(
    db,
    *,
    items,
    raw_rows=None,
    detail_rows=None,
    summary_rows=None,
    metrics=None,
    file_date="0701",
    analysis_date="2026-07-01",
):
    four_class = {
        "analysis_date": analysis_date,
        "total": len(items),
        "hit_count": sum(1 for it in items if str(it.get("status", "")).startswith("已触发")),
        "warn_count": sum(1 for it in items if it.get("status") == "预警"),
        "items": items,
        "summary": {},
    }
    rec = ZaigongRecord(
        source_filename="test.xlsx",
        file_date=file_date,
        summary_data=json.dumps(summary_rows or [], ensure_ascii=False),
        metrics_data=json.dumps(metrics or {}, ensure_ascii=False),
        detail_data=json.dumps(detail_rows or [], ensure_ascii=False),
        four_class_warnings=json.dumps(four_class, ensure_ascii=False),
        raw_data=json.dumps(raw_rows or [], ensure_ascii=False),
        uploaded_at=datetime.now(),
    )
    db.add(rec)
    db.commit()
    db.refresh(rec)
    return rec


def _item(**overrides):
    base = {
        "id": 1,
        "status": "已触发",
        "type": "列账不及时",
        "code": "GC001",
        "name": "项目A",
        "major": "5G",
        "acceptType": "一次验收",
        "manager": "张三",
        "keyDate": "2026-06-01",
        "keyDateLabel": "初验批复日期",
        "deadline": "",
        "projectStatus": "在建",
        "daysLabel": "逾期10天",
        "suggestion": "催收货",
    }
    base.update(overrides)
    return base


def _load_workbook(resp: httpx.Response):
    return openpyxl.load_workbook(io.BytesIO(resp.content))


def _find_data_row(ws, code: str):
    """在导出表中找工程编码所在数据行（跳过标题/分组行）。"""
    for row in ws.iter_rows(min_row=5, values_only=False):
        # D 列 = 工程编码
        if row[3].value == code:
            return row
    return None


class TestFourClassExportRouter:
    def test_formula_neutralizer_preserves_plain_values(self):
        assert analysis_mod.neutralize_spreadsheet_formula("普通项目") == "普通项目"
        assert analysis_mod.neutralize_spreadsheet_formula(123) == 123

    @pytest.mark.parametrize("value", [
        "=HYPERLINK(\"https://evil.test\",\"打开\")",
        "+SUM(1,1)",
        "-1+2",
        "@SUM(1,1)",
        "  =CMD()",
        "\t=CMD()",
    ])
    def test_formula_neutralizer_prefixes_dangerous_text(self, value):
        assert analysis_mod.neutralize_spreadsheet_formula(value).startswith("'")

    @pytest.mark.asyncio
    async def test_export_header_includes_construction_unit(self, test_db):
        Session = test_db
        db = Session()
        try:
            rec = _seed_record(db, items=[_item(constructionUnit="一公司")])
            rid = rec.id
        finally:
            db.close()

        async with build_client() as c:
            r = await c.get(f"/api/zaigong/four-class-warnings/{rid}/export")

        assert r.status_code == 200
        wb = _load_workbook(r)
        ws = wb.active
        headers = [cell.value for cell in ws[4]]
        assert "施工单位" in headers
        assert headers.index("施工单位") == 12  # 第 13 列（0-based 12）
        assert len([h for h in headers if h]) == 15

    @pytest.mark.asyncio
    async def test_export_uses_item_construction_unit(self, test_db):
        Session = test_db
        db = Session()
        try:
            rec = _seed_record(
                db,
                items=[_item(code="GC001", constructionUnit="一公司")],
                raw_rows=[{"工程编码": "GC001", "施工单位": "不该用这个"}],
            )
            rid = rec.id
        finally:
            db.close()

        async with build_client() as c:
            r = await c.get(f"/api/zaigong/four-class-warnings/{rid}/export")

        assert r.status_code == 200
        ws = _load_workbook(r).active
        row = _find_data_row(ws, "GC001")
        assert row is not None
        assert row[12].value == "一公司"

    @pytest.mark.asyncio
    async def test_export_neutralizes_formula_cells(self, test_db):
        Session = test_db
        db = Session()
        try:
            rec = _seed_record(
                db,
                items=[_item(code="GC-FORMULA", name='=HYPERLINK("https://evil.test","打开")')],
            )
            rid = rec.id
        finally:
            db.close()

        async with build_client() as c:
            r = await c.get(f"/api/zaigong/four-class-warnings/{rid}/export")

        assert r.status_code == 200
        ws = _load_workbook(r).active
        row = _find_data_row(ws, "GC-FORMULA")
        assert row is not None
        assert row[4].data_type != "f"
        assert row[4].value.startswith("'=")

    @pytest.mark.asyncio
    async def test_export_fallback_raw_data_when_unit_missing(self, test_db):
        """旧预警 JSON 无 constructionUnit 时，从 raw_data 按工程编码回查。"""
        Session = test_db
        db = Session()
        try:
            # 刻意不写 constructionUnit，模拟旧记录
            old_item = _item(code="GC009")
            old_item.pop("constructionUnit", None)
            rec = _seed_record(
                db,
                items=[old_item],
                raw_rows=[
                    {"工程编码": "GC009", "施工单位": "  回查施工单位  "},
                    {"工程编码": "GC010", "施工单位": "其他"},
                ],
            )
            rid = rec.id
        finally:
            db.close()

        async with build_client() as c:
            r = await c.get(f"/api/zaigong/four-class-warnings/{rid}/export")

        assert r.status_code == 200
        ws = _load_workbook(r).active
        row = _find_data_row(ws, "GC009")
        assert row is not None
        assert row[12].value == "回查施工单位"

    @pytest.mark.asyncio
    async def test_export_filename_contains_date_suffix(self, test_db):
        Session = test_db
        db = Session()
        try:
            rec = _seed_record(db, items=[_item()], file_date="0701")
            rid = rec.id
        finally:
            db.close()

        async with build_client() as c:
            r = await c.get(f"/api/zaigong/four-class-warnings/{rid}/export")

        assert r.status_code == 200
        cd = r.headers.get("content-disposition", "")
        # filename*=UTF-8''...
        m = re.search(r"filename\*=UTF-8''(.+)", cd, re.I)
        assert m, f"Content-Disposition 异常: {cd}"
        filename = unquote(m.group(1))
        today = datetime.now().strftime("%Y%m%d")
        assert filename.startswith("四类工程预警清单_0701_")
        assert today in filename
        assert filename.endswith(".xlsx")

    @pytest.mark.asyncio
    async def test_export_404_when_no_record(self, test_db):
        async with build_client() as c:
            r = await c.get("/api/zaigong/four-class-warnings/99999/export")
        assert r.status_code == 404
        assert r.json()["success"] is False


class TestAllManagerDetailsExportRouter:
    @pytest.mark.asyncio
    async def test_single_manager_export_keeps_original_sheet_layout(self, test_db):
        Session = test_db
        db = Session()
        try:
            rec = _seed_record(
                db,
                items=[],
                detail_rows=[
                    {
                        "工程名称": "项目A",
                        "工程管理员": "张三",
                        "施工单位": "施工一队",
                        "在建工程期末余额": 10,
                    },
                ],
                summary_rows=[{"工程管理员": "张三"}],
            )
            rid = rec.id
        finally:
            db.close()

        async with build_client() as c:
            r = await c.get(
                f"/api/zaigong/manager-details/{rid}/export",
                params={"manager": "张三"},
            )

        assert r.status_code == 200
        wb = _load_workbook(r)
        assert wb.sheetnames == ["张三工程明细"]
        assert [cell.value for cell in wb.active[10]][:4] == ["编号", "工程名称", "施工单位", "结转额"]
        assert wb.active["C11"].value == "施工一队"

    @pytest.mark.asyncio
    async def test_export_starts_with_all_projects_summary_sheet(self, test_db):
        detail_rows = [
            {
                "工程名称": "项目A",
                "工程管理员": "张三",
                "施工单位": "施工一队",
                "结转额": 20,
                "在建工程期末余额": 100,
                "本年累计资本性支出": 10000,
            },
            {
                "工程名称": "项目B",
                "工程管理员": "李四",
                "constructionUnit": "施工二队",
                "结转额": 10,
                "在建工程期末余额": 50,
                "本年累计资本性支出": 20000,
            },
        ]
        Session = test_db
        db = Session()
        try:
            rec = _seed_record(
                db,
                items=[],
                detail_rows=detail_rows,
                summary_rows=[
                    {"工程管理员": "张三"},
                    {"工程管理员": "李四"},
                    {"工程管理员": "合计", "转固率": 0.7317},
                ],
            )
            rid = rec.id
        finally:
            db.close()

        async with build_client() as c:
            r = await c.get(f"/api/zaigong/manager-details/{rid}/export-all")

        assert r.status_code == 200
        wb = _load_workbook(r)
        assert wb.sheetnames == ["全部工程汇总", "张三工程明细", "李四工程明细"]

        assert [cell.value for cell in wb["张三工程明细"][10]][:4] == [
            "编号", "工程名称", "施工单位", "结转额",
        ]
        assert wb["张三工程明细"]["C11"].value == "施工一队"
        assert wb["李四工程明细"]["C11"].value == "施工二队"

        ws = wb["全部工程汇总"]
        headers = [cell.value for cell in ws[10]]
        assert headers[:4] == ["编号", "工程名称", "工程管理员", "施工单位"]
        assert ws["B11"].value == "项目A"
        assert ws["C11"].value == "张三"
        assert ws["D11"].value == "施工一队"
        assert ws["B12"].value == "项目B"
        assert ws["C12"].value == "李四"
        assert ws["D12"].value == "施工二队"
        assert ws["C7"].value == pytest.approx(0.7317)
        assert ws["J13"].value == pytest.approx(0.7317)
        assert ws.auto_filter.ref == "A10:J12"

    @pytest.mark.asyncio
    async def test_summary_rate_falls_back_to_saved_core_metric(self, test_db):
        Session = test_db
        db = Session()
        try:
            rec = _seed_record(
                db,
                items=[],
                detail_rows=[{"工程名称": "项目A", "工程管理员": "张三", "在建工程期末余额": 10}],
                summary_rows=[{"工程管理员": "张三"}],
                metrics={"total_rate": 0.625},
            )
            rid = rec.id
        finally:
            db.close()

        async with build_client() as c:
            r = await c.get(f"/api/zaigong/manager-details/{rid}/export-all")

        assert r.status_code == 200
        ws = _load_workbook(r)["全部工程汇总"]
        assert ws["C7"].value == pytest.approx(0.625)
        assert ws["J12"].value == pytest.approx(0.625)

    @pytest.mark.asyncio
    async def test_summary_sheet_contains_all_projects_even_without_summary_manager(self, test_db):
        Session = test_db
        db = Session()
        try:
            rec = _seed_record(
                db,
                items=[],
                detail_rows=[{"工程名称": "未分配项目", "在建工程期末余额": 1}],
            )
            rid = rec.id
        finally:
            db.close()

        async with build_client() as c:
            r = await c.get(f"/api/zaigong/manager-details/{rid}/export-all")

        assert r.status_code == 200
        wb = _load_workbook(r)
        assert wb.sheetnames == ["全部工程汇总"]
        assert wb["全部工程汇总"]["C11"].value == "未分配"
